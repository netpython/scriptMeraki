#!/usr/bin/env python3
"""Rapport organisationnel Meraki : inventaire, santé et qualité des données."""
from __future__ import annotations
import sys
from collections import Counter
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from common import *
log = setup('organization_compliance')
args = base_parser('Conformité organisation Meraki').parse_args()
cfg = config(args.config)
api = dashboard()
oid = org_id(cfg)
names = network_name_map(api, oid)
statuses = {x['serial']: x for x in api.organizations.getOrganizationDevicesStatuses(oid, total_pages='all')}
rows = []
for d in api.organizations.getOrganizationDevices(oid, total_pages='all'):
    s = statuses.get(d['serial'], {})
    checks = {'name': bool(d.get('name')), 'network': bool(d.get('networkId')), 'tags': bool(d.get('tags')), 'lan_ip': bool(d.get('lanIp')), 'online': s.get('status') == 'online'}
    score = round(100 * sum(checks.values()) / len(checks))
    rows.append({'network': names.get(d.get('networkId'), ''), 'device': d.get('name'), 'serial': d['serial'], 'model': d.get('model'), 'status': s.get('status'), 'tags': ','.join(d.get('tags', [])), **{f'check_{k}': 'PASS' if v else 'FAIL' for k, v in checks.items()}, 'score': score, 'compliance': 'OK' if score == 100 else 'WARNING' if score >= 60 else 'CRITICAL'})
wb = Workbook()
ws = wb.active
ws.title = 'Audit'
headers = list(rows[0]) if rows else ['result']
ws.append(headers)
for r in rows:
    ws.append([r.get(h, '') for h in headers])
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='1F4E78')
ws.freeze_panes = 'A2'
summary = wb.create_sheet('Summary')
summary.append(['Status', 'Count'])
for k, v in Counter((r['compliance'] for r in rows)).items():
    summary.append([k, v])
path = OUTPUT / f'organization_compliance_{stamp()}.xlsx'
wb.save(path)
log.info('%s', path)
